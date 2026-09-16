import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Pt
from paths import OUTLINE_DOCX, OUTLINE_PDF  # noqa: E402

path = OUTLINE_DOCX
doc = Document(path)
print("段落总数:", len(doc.paragraphs))
print()
rows = []
for p in doc.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
    bold = any(r.bold for r in p.runs if r.bold)
    sz = max(sizes) if sizes else 0
    rows.append((sz, bold, t))
print("非空段落:", len(rows))
print()
print("=== 前 80 段（字号 | 粗 | 内容） ===")
for sz, b, t in rows[:80]:
    mark = "B" if b else " "
    print(f"{sz:>5} |{mark}| {t[:70]}")
